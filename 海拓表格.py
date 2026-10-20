import csv, os, tkinter as tk
from tkinter import ttk, filedialog, messagebox
PAGE_SIZE = 500

class TableApp:
    def __init__(self,root):
        self.root=root; root.title("海拓表格 0.1"); root.geometry("1200x760")
        self.headers=[];self.rows=[];self.filtered=None;self.page=0;self.path=None;self.editing=None
        self.status=tk.StringVar(value="就绪");self.search_var=tk.StringVar()
        top=ttk.Frame(root,padding=6);top.pack(fill="x")
        for name,action in [("打开文件",self.open_file),("保存",self.save),("另存为",self.save_as),("新增行",self.add_row),("删除选中行",self.delete_rows),("查找",self.search),("清除筛选",self.clear_filter),("按当前列去重",self.deduplicate),("上一页",self.prev_page),("下一页",self.next_page)]:
            ttk.Button(top,text=name,command=action).pack(side="left",padx=2)
        ttk.Entry(top,textvariable=self.search_var,width=20).pack(side="left",padx=4)
        self.field_var=tk.StringVar(value="所有列");self.field_combo=ttk.Combobox(root,textvariable=self.field_var,state="readonly",width=24)
        self.field_combo.pack(anchor="w",padx=9,pady=(2,3))
        frame=ttk.Frame(root);frame.pack(fill="both",expand=True,padx=8,pady=4)
        self.tree=ttk.Treeview(frame,show="headings",selectmode="extended")
        sy=ttk.Scrollbar(frame,orient="vertical",command=self.tree.yview);sx=ttk.Scrollbar(frame,orient="horizontal",command=self.tree.xview)
        self.tree.configure(yscrollcommand=sy.set,xscrollcommand=sx.set)
        self.tree.grid(row=0,column=0,sticky="nsew");sy.grid(row=0,column=1,sticky="ns");sx.grid(row=1,column=0,sticky="ew")
        frame.rowconfigure(0,weight=1);frame.columnconfigure(0,weight=1)
        self.tree.bind("<Double-1>",self.start_edit)
        ttk.Label(root,textvariable=self.status,padding=5).pack(fill="x")
        self.new_table()

    def new_table(self):
        self.headers=["A","B","C","D","E"];self.rows=[[""]*5 for _ in range(30)]
        self.filtered=None;self.page=0;self.configure_columns();self.render()

    def configure_columns(self):
        cols=[f"c{i}" for i in range(len(self.headers))]
        self.tree.configure(columns=cols)
        for i,c in enumerate(cols):
            self.tree.heading(c,text=self.headers[i]);self.tree.column(c,width=165,minwidth=70,stretch=False)
        self.field_combo["values"]=["所有列"]+self.headers;self.field_var.set("所有列")

    def indices(self):return self.filtered if self.filtered is not None else range(len(self.rows))

    def render(self):
        self.tree.delete(*self.tree.get_children())
        ids=self.indices();start=self.page*PAGE_SIZE
        for idx in list(ids)[start:start+PAGE_SIZE]:self.tree.insert("","end",iid=str(idx),values=self.rows[idx])
        total=len(ids)
        self.status.set(f"共 {len(self.rows):,} 行 | 当前筛选 {total:,} 行 | 第 {self.page+1}/{max(1,(total+PAGE_SIZE-1)//PAGE_SIZE)} 页 | 双击单元格编辑")

    def open_file(self):
        path=filedialog.askopenfilename(filetypes=[("表格文件","*.csv *.tsv *.txt *.xlsx"),("所有文件","*.*")])
        if not path:return
        try:
            ext=os.path.splitext(path)[1].lower()
            if ext==".xlsx":
                from openpyxl import load_workbook
                wb=load_workbook(path,read_only=True,data_only=False);ws=wb.active
                data=[[str(v) if v is not None else "" for v in row] for row in ws.iter_rows(values_only=True)]
                wb.close()
            else:
                delimiter="\t" if ext in (".tsv",".txt") else ","
                data=None
                for enc in ("utf-8-sig","utf-8","gb18030"):
                    try:
                        with open(path,newline="",encoding=enc) as f:data=list(csv.reader(f,delimiter=delimiter))
                        break
                    except UnicodeDecodeError:pass
                if data is None:raise ValueError("无法识别文件编码")
            if not data:raise ValueError("文件为空")
            width=max(map(len,data))
            self.headers=[str(x) or f"列{i+1}" for i,x in enumerate(data[0])]+[f"列{i+1}" for i in range(len(data[0]),width)]
            self.rows=[row+[""]*(width-len(row)) for row in data[1:]]
            self.path=path;self.filtered=None;self.page=0;self.configure_columns();self.render()
            self.root.title("海拓表格 0.1 — "+os.path.basename(path))
        except Exception as e:messagebox.showerror("打开失败",str(e))

    def save(self):
        if not self.path:self.save_as();return
        self.write(self.path)

    def save_as(self):
        path=filedialog.asksaveasfilename(defaultextension=".csv",filetypes=[("CSV 文件","*.csv"),("Excel 文件","*.xlsx"),("制表符文本","*.tsv")])
        if path:self.write(path)

    def write(self,path):
        try:
            if path.lower().endswith(".xlsx"):
                from openpyxl import Workbook
                wb=Workbook(write_only=True);ws=wb.create_sheet("表格");ws.append(self.headers)
                for row in self.rows:ws.append(row)
                wb.save(path)
            else:
                delimiter="\t" if path.lower().endswith(".tsv") else ","
                with open(path,"w",newline="",encoding="utf-8-sig") as f:
                    writer=csv.writer(f,delimiter=delimiter);writer.writerow(self.headers);writer.writerows(self.rows)
            self.path=path;self.status.set("已保存："+path)
        except Exception as e:messagebox.showerror("保存失败",str(e))

    def start_edit(self,event):
        iid=self.tree.identify_row(event.y);column=self.tree.identify_column(event.x)
        if not iid or not column:return
        ci=int(column[1:])-1;box=self.tree.bbox(iid,column)
        if not box:return
        if self.editing:self.editing.destroy()
        x,y,w,h=box;editor=ttk.Entry(self.tree);editor.insert(0,self.rows[int(iid)][ci])
        editor.place(x=x,y=y,width=w,height=h);editor.focus_set();editor.select_range(0,"end");self.editing=editor
        def finish(_=None):
            if self.editing is not editor:return
            self.rows[int(iid)][ci]=editor.get();editor.destroy();self.editing=None
            if self.tree.exists(iid):self.tree.set(iid,f"c{ci}",self.rows[int(iid)][ci])
        editor.bind("<Return>",finish);editor.bind("<FocusOut>",finish)
        editor.bind("<Escape>",lambda _: (editor.destroy(),setattr(self,"editing",None)))

    def add_row(self):
        self.rows.append([""]*len(self.headers));self.filtered=None;self.page=(len(self.rows)-1)//PAGE_SIZE;self.render()

    def delete_rows(self):
        chosen={int(i) for i in self.tree.selection()}
        if not chosen:return
        if not messagebox.askyesno("确认删除",f"删除选中的 {len(chosen)} 行吗？"):return
        self.rows=[r for i,r in enumerate(self.rows) if i not in chosen]
        self.filtered=None;self.page=0;self.render()

    def search(self):
        keyword=self.search_var.get().casefold()
        if not keyword:self.clear_filter();return
        field=self.field_var.get();ci=self.headers.index(field) if field!="所有列" and field in self.headers else None
        self.filtered=[i for i,row in enumerate(self.rows) if any(keyword in str(v).casefold() for v in ([row[ci]] if ci is not None else row))]
        self.page=0;self.render()

    def clear_filter(self):self.filtered=None;self.page=0;self.render()

    def deduplicate(self):
        field=self.field_var.get()
        if field=="所有列":messagebox.showinfo("提示","请先在下拉框中选择用于去重的列");return
        ci=self.headers.index(field)
        if not messagebox.askyesno("确认去重",f"按「{field}」列去重，保留首次出现的行（空值不去重）？"):return
        found=set();kept=[];removed=0
        for row in self.rows:
            key=row[ci].strip().casefold()
            if key and key in found:removed+=1;continue
            if key:found.add(key)
            kept.append(row)
        self.rows=kept;self.filtered=None;self.page=0;self.render()
        messagebox.showinfo("完成",f"删除 {removed:,} 行重复记录")

    def prev_page(self):self.page=max(0,self.page-1);self.render()
    def next_page(self):
        if (self.page+1)*PAGE_SIZE<len(self.indices()):self.page+=1;self.render()

if __name__=="__main__":
    root=tk.Tk();style=ttk.Style(root)
    if "vista" in style.theme_names():style.theme_use("vista")
    app=TableApp(root);root.mainloop()
